"""Ekspor laporan ke PDF (reportlab) dan xlsx (openpyxl): FR-LAP-04, NFR-PRF-03, NFR-SEC-05.

Baris diambil dari `laporan.laporan_transaksi` / `laporan.laporan_tagihan` tanpa halaman, jadi isi
ekspor selalu sama dengan JSON untuk filter yang sama (satu sumber filter).

- xlsx: setiap sel teks dipaksa bertipe string (`data_type = "s"`), sehingga data pengguna yang
  diawali `=` (nama, judul) tidak pernah tersimpan sebagai rumus (NFR-SEC-05). Tanggal dan nominal
  disimpan sebagai sel tanggal/angka dengan format tampilan, agar tetap bisa dijumlah.
- PDF: font DejaVu Sans tertanam (`app/aset/font/`, lisensi di `LICENSE-DejaVu.txt`) agar
  karakter di luar Latin-1 tidak hilang; teks dari data di-escape sebelum masuk `Paragraph`
  (bukan markup).
- Waktu cetak dari jam aplikasi dalam WIB (`sekarang_wib`, K-07).
"""

import io
from dataclasses import dataclass
from datetime import date
from functools import cache
from pathlib import Path
from typing import Literal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from sqlalchemy.orm import Session

from app.core.format import format_rupiah, format_tanggal
from app.core.waktu import hari_ini_wib, sekarang_wib
from app.models.status import (
    LABEL_CARA_PENYELESAIAN,
    LABEL_JENIS_TAGIHAN,
    LABEL_STATUS_ITEM,
    LABEL_STATUS_TAGIHAN,
    LABEL_TERLAMBAT,
    CaraPenyelesaian,
    JenisTagihan,
    StatusItem,
    StatusTagihan,
)
from app.services import laporan
from app.services.laporan import FilterTagihan, FilterTransaksi, StatusLaporan

MIME_PDF = "application/pdf"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
Format = Literal["pdf", "xlsx"]

_FOLDER_FONT = Path(__file__).resolve().parent.parent / "aset" / "font"
_FONT, _FONT_TEBAL = "DejaVuSans", "DejaVuSans-Bold"
_FORMAT_TANGGAL_XLSX = "DD/MM/YYYY"
_FORMAT_RUPIAH_XLSX = '"Rp"#,##0'
_LABEL_STATUS_LAPORAN = {
    StatusLaporan.DIPINJAM: "Dipinjam",
    StatusLaporan.TERLAMBAT: LABEL_TERLAMBAT,
    StatusLaporan.DIKEMBALIKAN: "Dikembalikan",
    StatusLaporan.HILANG: "Hilang",
    StatusLaporan.RUSAK: "Rusak",
}

Jenis = Literal["teks", "tanggal", "rupiah"]


@dataclass(frozen=True)
class BerkasEkspor:
    isi: bytes
    media_type: str
    nama_berkas: str


@dataclass(frozen=True)
class _Kolom:
    judul: str
    jenis: Jenis
    lebar_mm: float
    bungkus: bool = False  # teks panjang dibungkus (Paragraph) di PDF


@dataclass(frozen=True)
class _Laporan:
    judul: str
    nama_dasar: str  # `laporan-<nama_dasar>-YYYYMMDD.<ext>`
    keterangan: list[str]  # filter aktif + waktu cetak
    kolom: list[_Kolom]
    baris: list[list[str | date | int | None]]
    total_nominal: int | None = None  # baris "Total" (laporan tagihan)


# --------------------------------------------------------------------------------------- font


@cache
def font_pdf() -> TTFont:
    """Daftarkan DejaVu Sans (regular & tebal) sekali; dibaca dari repo, tidak diunduh."""
    biasa = TTFont(_FONT, str(_FOLDER_FONT / "DejaVuSans.ttf"))
    pdfmetrics.registerFont(biasa)
    pdfmetrics.registerFont(TTFont(_FONT_TEBAL, str(_FOLDER_FONT / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFontFamily(_FONT, normal=_FONT, bold=_FONT_TEBAL)
    return biasa


# --------------------------------------------------------------------------------------- isi


def _rentang(dari: date | None, sampai: date | None) -> str:
    if dari and sampai:
        return f"{format_tanggal(dari)} s.d. {format_tanggal(sampai)}"
    if dari:
        return f"mulai {format_tanggal(dari)}"
    if sampai:
        return f"sampai {format_tanggal(sampai)}"
    return "Semua"


def _label(peta: dict, kode: str | None) -> str:
    return "Semua" if kode is None else peta[type(next(iter(peta)))(kode)]


def _dicetak() -> str:
    return f"Dicetak: {sekarang_wib():%d/%m/%Y %H:%M} WIB"


def _isi_transaksi(db: Session, f: FilterTransaksi) -> _Laporan:
    baris, total = laporan.laporan_transaksi(db, f)
    return _Laporan(
        judul="Laporan Transaksi Peminjaman dan Pengembalian",
        nama_dasar="transaksi",
        keterangan=[
            f"Rentang tanggal pinjam: {_rentang(f.dari, f.sampai)}",
            f"Status: {'Semua' if f.status is None else _LABEL_STATUS_LAPORAN[f.status]}",
            f"Jumlah baris: {total}",
            _dicetak(),
        ],
        kolom=[
            _Kolom("Kode Anggota", "teks", 24),
            _Kolom("Nama Anggota", "teks", 52, bungkus=True),
            _Kolom("Judul", "teks", 74, bungkus=True),
            _Kolom("Kode Eksemplar", "teks", 26),
            _Kolom("Tanggal Pinjam", "tanggal", 23),
            _Kolom("Jatuh Tempo", "tanggal", 23),
            _Kolom("Tanggal Kembali", "tanggal", 23),
            _Kolom("Status", "teks", 24),
        ],
        baris=[
            [
                b.anggota_kode,
                b.anggota_nama,
                b.judul,
                b.kode_eksemplar,
                b.tanggal_pinjam,
                b.jatuh_tempo,
                b.tanggal_kembali,
                LABEL_TERLAMBAT if b.terlambat else LABEL_STATUS_ITEM[StatusItem(b.status)],
            ]
            for b in baris
        ],
    )


def _isi_tagihan(db: Session, f: FilterTagihan) -> _Laporan:
    h = laporan.laporan_tagihan(db, f)
    return _Laporan(
        judul="Laporan Denda dan Penggantian",
        nama_dasar="tagihan",
        keterangan=[
            f"Rentang tanggal dibentuk: {_rentang(f.dari, f.sampai)}",
            f"Jenis: {_label(LABEL_JENIS_TAGIHAN, f.jenis)}",
            f"Status: {_label(LABEL_STATUS_TAGIHAN, f.status)}",
            f"Cara Penyelesaian: {_label(LABEL_CARA_PENYELESAIAN, f.cara)}",
            f"Jumlah tagihan: {h.jumlah} · Total nominal: {format_rupiah(h.total_nominal)}",
            _dicetak(),
        ],
        kolom=[
            _Kolom("Tanggal Dibentuk", "tanggal", 21),
            _Kolom("Kode Anggota", "teks", 22),
            _Kolom("Nama Anggota", "teks", 34, bungkus=True),
            _Kolom("Judul", "teks", 45, bungkus=True),
            _Kolom("Kode Eksemplar", "teks", 22),
            _Kolom("Jenis", "teks", 21),
            _Kolom("Nominal", "rupiah", 24),
            _Kolom("Status", "teks", 19),
            _Kolom("Cara Penyelesaian", "teks", 22),
            _Kolom("Tanggal Penyelesaian", "tanggal", 21),
            _Kolom("Admin Pengonfirmasi", "teks", 22, bungkus=True),
        ],
        baris=[
            [
                b.tanggal_dibentuk,
                b.anggota_kode,
                b.anggota_nama,
                b.judul,
                b.kode_eksemplar,
                LABEL_JENIS_TAGIHAN[JenisTagihan(b.jenis)],
                b.nominal,
                LABEL_STATUS_TAGIHAN[StatusTagihan(b.status)],
                LABEL_CARA_PENYELESAIAN[CaraPenyelesaian(b.cara_penyelesaian)]
                if b.cara_penyelesaian
                else None,
                b.tanggal_penyelesaian,
                b.admin_pengonfirmasi,
            ]
            for b in h.baris
        ],
        total_nominal=h.total_nominal,
    )


# --------------------------------------------------------------------------------------- xlsx


def _sel_teks(ws, baris: int, kolom: int, nilai: str):
    """NFR-SEC-05: paksa tipe string; openpyxl menganggap teks berawalan `=` sebagai rumus."""
    sel = ws.cell(row=baris, column=kolom)
    sel.value = nilai
    sel.data_type = "s"
    return sel


def _xlsx(lap: _Laporan) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = lap.nama_dasar.capitalize()
    _sel_teks(ws, 1, 1, lap.judul).font = Font(bold=True, size=13)
    for n, teks in enumerate(lap.keterangan, start=2):
        _sel_teks(ws, n, 1, teks)
    kepala = len(lap.keterangan) + 3  # satu baris kosong sebelum tabel
    for k, kol in enumerate(lap.kolom, start=1):
        sel = _sel_teks(ws, kepala, k, kol.judul)
        sel.font = Font(bold=True)
        sel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(k)].width = kol.lebar_mm / 2.2
    r = kepala
    for r, isi in enumerate(lap.baris, start=kepala + 1):
        for k, (kol, nilai) in enumerate(zip(lap.kolom, isi, strict=True), start=1):
            if nilai is None:
                continue
            if kol.jenis == "teks":
                _sel_teks(ws, r, k, nilai)
                continue
            sel = ws.cell(row=r, column=k, value=nilai)
            sel.number_format = (
                _FORMAT_TANGGAL_XLSX if kol.jenis == "tanggal" else _FORMAT_RUPIAH_XLSX
            )
    if lap.total_nominal is not None:
        r += 1
        _sel_teks(ws, r, 1, "Total").font = Font(bold=True)
        k_nominal = next(k for k, kol in enumerate(lap.kolom, start=1) if kol.jenis == "rupiah")
        sel = ws.cell(row=r, column=k_nominal, value=lap.total_nominal)
        sel.number_format = _FORMAT_RUPIAH_XLSX
        sel.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=kepala + 1, column=1)
    keluaran = io.BytesIO()
    wb.save(keluaran)
    return keluaran.getvalue()


# --------------------------------------------------------------------------------------- PDF


def _teks_pdf(nilai, kol: _Kolom, gaya: ParagraphStyle):
    if nilai is None:
        return ""
    if kol.jenis == "tanggal":
        return format_tanggal(nilai)
    if kol.jenis == "rupiah":
        return format_rupiah(nilai)
    # Sel string biasa digambar apa adanya; Paragraph menafsirkan markup, jadi wajib di-escape.
    return Paragraph(escape(nilai), gaya) if kol.bungkus else nilai


def _nomor_halaman(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont(_FONT, 7)
    canvas.drawRightString(doc.pagesize[0] - 12 * mm, 7 * mm, f"Halaman {doc.page}")
    canvas.restoreState()


def _pdf(lap: _Laporan) -> bytes:
    font_pdf()
    sel = ParagraphStyle("sel", fontName=_FONT, fontSize=7, leading=8.5)
    kepala_sel = ParagraphStyle("kepala_sel", parent=sel, fontName=_FONT_TEBAL)
    judul = ParagraphStyle("judul", fontName=_FONT_TEBAL, fontSize=13, leading=16, spaceAfter=4)
    ket = ParagraphStyle("ket", fontName=_FONT, fontSize=8, leading=10)

    data = [[Paragraph(escape(k.judul), kepala_sel) for k in lap.kolom]]
    data += [
        [_teks_pdf(v, k, sel) for k, v in zip(lap.kolom, isi, strict=True)] for isi in lap.baris
    ]
    gaya = [
        ("FONT", (0, 0), (-1, -1), _FONT, 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    for k, kol in enumerate(lap.kolom):
        if kol.jenis == "rupiah":
            gaya.append(("ALIGN", (k, 1), (k, -1), "RIGHT"))
    if lap.total_nominal is not None:
        k_nominal = next(k for k, kol in enumerate(lap.kolom) if kol.jenis == "rupiah")
        total = [""] * len(lap.kolom)
        total[0], total[k_nominal] = "Total", format_rupiah(lap.total_nominal)
        data.append(total)
        gaya.append(("FONT", (0, -1), (-1, -1), _FONT_TEBAL, 7))

    tabel = LongTable(data, colWidths=[k.lebar_mm * mm for k in lap.kolom], repeatRows=1)
    tabel.setStyle(TableStyle(gaya))
    keluaran = io.BytesIO()
    doc = SimpleDocTemplate(
        keluaran,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=lap.judul,
    )
    cerita = [Paragraph(escape(lap.judul), judul)]
    cerita += [Paragraph(escape(t), ket) for t in lap.keterangan]
    cerita += [Spacer(1, 4 * mm), tabel]
    doc.build(cerita, onFirstPage=_nomor_halaman, onLaterPages=_nomor_halaman)
    return keluaran.getvalue()


# --------------------------------------------------------------------------------------- publik


def _berkas(lap: _Laporan, format: Format) -> BerkasEkspor:
    nama = f"laporan-{lap.nama_dasar}-{hari_ini_wib():%Y%m%d}.{format}"
    if format == "pdf":
        return BerkasEkspor(_pdf(lap), MIME_PDF, nama)
    return BerkasEkspor(_xlsx(lap), MIME_XLSX, nama)


def ekspor_transaksi(db: Session, f: FilterTransaksi, format: Format) -> BerkasEkspor:
    """FR-LAP-04: semua baris laporan transaksi (FR-LAP-02) sesuai filter aktif."""
    return _berkas(_isi_transaksi(db, f), format)


def ekspor_tagihan(db: Session, f: FilterTagihan, format: Format) -> BerkasEkspor:
    """FR-LAP-04: semua baris laporan tagihan (FR-LAP-03) + total sesuai filter aktif."""
    return _berkas(_isi_tagihan(db, f), format)
