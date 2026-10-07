"""WP 5.2.3 + 5.3.13 — ekspor laporan PDF & xlsx: FR-LAP-04, NFR-PRF-03, NFR-SEC-05, OQ-37..41.

Isi berkas dibaca balik (openpyxl untuk xlsx, pypdf untuk teks PDF) dan dibandingkan dengan JSON
tanpa halaman untuk filter yang sama (satu sumber filter).
"""

import io
import re
import time
import zipfile
from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, insert, select, text
from sqlalchemy.orm import Session

from app.models import Eksemplar, ItemTransaksi, Tagihan, TransaksiPeminjaman
from app.services import ekspor, laporan
from app.services.laporan import FilterTransaksi, StatusLaporan
from tests import pabrik
from tests.test_laporan import _item, _tagihan

HARI_INI = date(2026, 11, 15)
API_TRX = "/api/v1/admin/laporan/transaksi"
API_TGH = "/api/v1/admin/laporan/tagihan"
KOLOM_TRX = [
    "Kode Anggota",
    "Nama Anggota",
    "Judul",
    "Kode Eksemplar",
    "Tanggal Pinjam",
    "Jatuh Tempo",
    "Tanggal Kembali",
    "Status",
]
KOLOM_TGH = [
    "Tanggal Dibentuk",
    "Kode Anggota",
    "Nama Anggota",
    "Judul",
    "Kode Eksemplar",
    "Jenis",
    "Nominal",
    "Status",
    "Cara Penyelesaian",
    "Tanggal Penyelesaian",
    "Admin Pengonfirmasi",
]
NON_LATIN = "Ольга Ἀθηνᾶ Đặng Thị Ngọc Şükrü Øster"


@pytest.fixture(autouse=True)
def jam(atur_waktu):
    atur_waktu(datetime(2026, 11, 15, 3, 0, tzinfo=UTC))  # 10:00 WIB
    return atur_waktu


@pytest.fixture(autouse=True)
def db_kosong(db: Session):
    for model in (ItemTransaksi, Tagihan):
        jumlah = db.scalar(select(func.count()).select_from(model))
        assert jumlah == 0, f"DB test harus kosong: {model.__tablename__} berisi {jumlah} baris"


# --------------------------------------------------------------------------- pembaca berkas


def _unduh(klien, path: str, **params):
    r = klien.get(path, params=params)
    assert r.status_code == 200, r.text
    return r


def _tabel_xlsx(isi: bytes, kolom: list[str]) -> tuple[list[list], list[list], object]:
    """(baris kepala sebelum tabel, baris data, lembar kerja). Kepala tabel dicari dari kolom."""
    ws = load_workbook(io.BytesIO(isi)).active
    semua = [[c for c in baris] for baris in ws.iter_rows()]
    i = next(n for n, b in enumerate(semua) if [c.value for c in b[: len(kolom)]] == kolom)
    data = []
    for b in semua[i + 1 :]:
        if b[0].value in (None, "Total"):
            break
        data.append(b)
    return semua[:i], data, ws


def _teks_pdf(isi: bytes) -> str:
    assert isi.startswith(b"%PDF")
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(isi)).pages)


def _kode_eks(db, item: ItemTransaksi) -> str:
    return db.get(Eksemplar, item.eksemplar_id).kode


@pytest.fixture
def data_trx(db: Session) -> dict[str, ItemTransaksi]:
    return {
        "dipinjam": _item(db, pinjam=HARI_INI - timedelta(days=5), judul="Masih Dipinjam"),
        "terlambat1": _item(db, pinjam=HARI_INI - timedelta(days=40), judul="Telat Satu"),
        "terlambat2": _item(db, pinjam=HARI_INI - timedelta(days=35), judul="Telat Dua"),
        "kembali": _item(
            db, pinjam=HARI_INI - timedelta(days=60), status="DIKEMBALIKAN", judul="Sudah Kembali"
        ),
        "hilang": _item(
            db, pinjam=HARI_INI - timedelta(days=50), status="HILANG", judul="Buku Hilang"
        ),
    }


# --------------------------------------------------------------------------- FR-LAP-04 xlsx


@pytest.mark.parametrize(
    "params",
    [
        {"status": "TERLAMBAT"},
        {"status": "DIPINJAM"},
        {
            "dari": (HARI_INI - timedelta(days=50)).isoformat(),
            "sampai": (HARI_INI - timedelta(days=35)).isoformat(),
        },
        {},
    ],
)
def test_FR_LAP_04_xlsx_transaksi_sama_dengan_json_tanpa_halaman(klien_admin, data_trx, params):
    json_ = _unduh(klien_admin, API_TRX, per_halaman=100, **params).json()
    berkas = _unduh(klien_admin, f"{API_TRX}/ekspor", format="xlsx", **params).content
    _, data, _ = _tabel_xlsx(berkas, KOLOM_TRX)

    def label(b):
        return "Terlambat" if b["terlambat"] else b["status"].title()

    assert [[c.value for c in baris[:4]] + [baris[7].value] for baris in data] == [
        [b["anggota_kode"], b["anggota_nama"], b["judul"], b["kode_eksemplar"], label(b)]
        for b in json_["data"]
    ]
    assert len(data) == json_["total"]


def test_FR_LAP_04_xlsx_kepala_memuat_filter_aktif_dan_waktu_cetak(klien_admin, data_trx):
    berkas = _unduh(
        klien_admin,
        f"{API_TRX}/ekspor",
        format="xlsx",
        status="TERLAMBAT",
        dari="2026-10-01",
        sampai="2026-10-31",
    ).content
    kepala, data, _ = _tabel_xlsx(berkas, KOLOM_TRX)
    teks = " ".join(str(c.value) for b in kepala for c in b if c.value is not None)
    assert "Laporan Transaksi Peminjaman dan Pengembalian" in teks
    assert "01/10/2026" in teks and "31/10/2026" in teks
    assert "Status: Terlambat" in teks
    assert "15/11/2026 10:00 WIB" in teks
    assert [b[2].value for b in data] == ["Telat Satu", "Telat Dua"]  # pinjam 06/10 & 11/10


def test_FR_LAP_04_xlsx_tipe_sel_tanggal_dan_nominal(klien_admin, db):
    _tagihan(db, dibentuk=date(2026, 10, 2), nominal=15_555, cara="TUNAI")
    _tagihan(db, dibentuk=date(2026, 10, 3), jenis="PENGGANTIAN", nominal=75_000)
    berkas = _unduh(klien_admin, f"{API_TGH}/ekspor", format="xlsx").content
    _, data, ws = _tabel_xlsx(berkas, KOLOM_TGH)
    assert len(data) == 2
    tgl, nominal = data[0][0], data[0][6]
    assert (
        tgl.is_date and tgl.value.date() == date(2026, 10, 2) and tgl.number_format == "DD/MM/YYYY"
    )
    assert isinstance(nominal.value, int) and nominal.value == 15_555
    assert "Rp" in nominal.number_format
    assert (data[0][5].value, data[0][7].value, data[0][8].value) == ("Denda", "Lunas", "Tunai")
    assert (data[1][5].value, data[1][7].value, data[1][8].value) == (
        "Penggantian",
        "Belum Lunas",
        None,
    )
    total = next(b for b in ws.iter_rows() if b[0].value == "Total")
    assert total[6].value == 15_555 + 75_000 and isinstance(total[6].value, int)
    assert "Rp" in total[6].number_format


def test_FR_LAP_04_xlsx_tagihan_mengikuti_filter_dan_total(klien_admin, db):
    _tagihan(db, dibentuk=date(2026, 10, 2), nominal=10_000)
    _tagihan(db, dibentuk=date(2026, 10, 3), nominal=20_000, cara="TRANSFER")
    _tagihan(db, dibentuk=date(2026, 10, 4), nominal=30_000, cara="TRANSFER")
    params = {"status": "LUNAS", "cara": "TRANSFER"}
    json_ = _unduh(klien_admin, API_TGH, per_halaman=100, **params).json()
    berkas = _unduh(klien_admin, f"{API_TGH}/ekspor", format="xlsx", **params).content
    kepala, data, ws = _tabel_xlsx(berkas, KOLOM_TGH)
    assert [b[4].value for b in data] == [b["kode_eksemplar"] for b in json_["data"]]
    total = next(b for b in ws.iter_rows() if b[0].value == "Total")
    assert total[6].value == json_["total_nominal"] == 50_000
    teks = " ".join(str(c.value) for b in kepala for c in b if c.value is not None)
    assert "Status: Lunas" in teks and "Cara Penyelesaian: Transfer" in teks
    assert "Jenis: Semua" in teks


# --------------------------------------------------------------------------- NFR-SEC-05 injeksi


def test_NFR_SEC_05_xlsx_teks_diawali_sama_dengan_tidak_menjadi_rumus(klien_admin, db):
    a = pabrik.anggota(db, nama="=1+1")
    _item(db, pinjam=HARI_INI - timedelta(days=3), judul='=HYPERLINK("http://x")', anggota=a)
    _tagihan(db, dibentuk=date(2026, 10, 2), judul="@SUM(A1)", anggota=a)
    kasus = (
        (API_TRX, KOLOM_TRX, {"=1+1", '=HYPERLINK("http://x")'}),
        (API_TGH, KOLOM_TGH, {"=1+1", "@SUM(A1)"}),
    )
    for path, kolom, harapan in kasus:
        berkas = _unduh(klien_admin, f"{path}/ekspor", format="xlsx").content
        _, data, _ = _tabel_xlsx(berkas, kolom)
        teks = [c for b in data for c in b if isinstance(c.value, str)]
        assert harapan <= {c.value for c in teks}, path  # literal, apa adanya
        assert all(c.data_type == "s" for c in teks), path
        with zipfile.ZipFile(io.BytesIO(berkas)) as z:
            for n in z.namelist():
                if n.startswith("xl/worksheets/sheet"):
                    assert "<f>" not in z.read(n).decode(), (path, n)  # tak ada sel rumus


# --------------------------------------------------------------------------- FR-LAP-04 PDF


def test_FR_LAP_04_pdf_transaksi_mengikuti_filter(klien_admin, db, data_trx):
    teks = _teks_pdf(
        _unduh(klien_admin, f"{API_TRX}/ekspor", format="pdf", status="TERLAMBAT").content
    )
    assert "Laporan Transaksi Peminjaman dan Pengembalian" in teks
    assert "Status: Terlambat" in teks
    assert "Dicetak: 15/11/2026 10:00 WIB" in teks
    for nama in ("terlambat1", "terlambat2"):
        assert _kode_eks(db, data_trx[nama]) in teks
    for nama in ("dipinjam", "kembali", "hilang"):
        assert _kode_eks(db, data_trx[nama]) not in teks
    assert "Terlambat" in teks
    assert (HARI_INI - timedelta(days=40)).strftime("%d/%m/%Y") in teks  # tanggal DD/MM/YYYY


def test_FR_LAP_04_pdf_tagihan_memuat_total_rupiah(klien_admin, db):
    _tagihan(db, dibentuk=date(2026, 10, 2), nominal=15_555)
    _tagihan(db, dibentuk=date(2026, 10, 3), nominal=1_000_000, jenis="PENGGANTIAN")
    _tagihan(db, dibentuk=date(2026, 10, 9), nominal=999, cara="TUNAI")
    teks = _teks_pdf(
        _unduh(klien_admin, f"{API_TGH}/ekspor", format="pdf", status="BELUM_LUNAS").content
    )
    assert "Laporan Denda dan Penggantian" in teks
    assert "Rp15.555" in teks and "Rp1.000.000" in teks and "Rp999" not in teks
    assert re.search(r"Total.*Rp1\.015\.555", teks.replace("\n", " "))
    assert "Status: Belum Lunas" in teks


def test_FR_LAP_04_pdf_karakter_non_latin_dan_markup_tidak_merusak(klien_admin, db):
    a = pabrik.anggota(db, nama=NON_LATIN)
    _item(db, pinjam=HARI_INI - timedelta(days=3), judul="<b>Judul & Tebal</b>", anggota=a)
    teks = _teks_pdf(_unduh(klien_admin, f"{API_TRX}/ekspor", format="pdf").content)
    rapat = re.sub(r"\s+", "", teks)
    assert re.sub(r"\s+", "", NON_LATIN) in rapat  # karakter tampil, tidak hilang/kotak
    assert "<b>Judul&Tebal</b>" in rapat  # markup ditulis apa adanya, tidak ditafsirkan


def test_FR_LAP_04_font_pdf_memuat_semua_glyph_non_latin():
    """Glyph tersedia di font tertanam (bukan kotak `.notdef`)."""
    font = ekspor.font_pdf()
    tanpa_glyph = [
        ch for ch in NON_LATIN if not ch.isspace() and ord(ch) not in font.face.charToGlyph
    ]
    assert tanpa_glyph == []


# --------------------------------------------------------------------------- respons berkas


@pytest.mark.parametrize(
    ("path", "fmt", "mime", "nama"),
    [
        (API_TRX, "pdf", "application/pdf", "laporan-transaksi-20261115.pdf"),
        (API_TRX, "xlsx", ekspor.MIME_XLSX, "laporan-transaksi-20261115.xlsx"),
        (API_TGH, "pdf", "application/pdf", "laporan-tagihan-20261115.pdf"),
        (API_TGH, "xlsx", ekspor.MIME_XLSX, "laporan-tagihan-20261115.xlsx"),
    ],
)
def test_FR_LAP_04_header_respons_berkas(klien_admin, path, fmt, mime, nama):
    r = _unduh(klien_admin, f"{path}/ekspor", format=fmt)
    assert r.headers["content-type"] == mime
    assert r.headers["content-disposition"] == f'attachment; filename="{nama}"'
    assert ekspor.MIME_XLSX == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_FR_LAP_04_laporan_kosong_tetap_berkas_sah(klien_admin):
    assert _teks_pdf(_unduh(klien_admin, f"{API_TGH}/ekspor", format="pdf").content)
    _, data, ws = _tabel_xlsx(
        _unduh(klien_admin, f"{API_TGH}/ekspor", format="xlsx").content, KOLOM_TGH
    )
    assert data == []
    assert next(b for b in ws.iter_rows() if b[0].value == "Total")[6].value == 0


@pytest.mark.parametrize("params", [{"format": "csv"}, {"format": "PDF"}, {}])
@pytest.mark.parametrize("path", [API_TRX, API_TGH])
def test_FR_LAP_04_format_tidak_dikenal_ditolak(klien_admin, path, params):
    r = klien_admin.get(f"{path}/ekspor", params=params)
    assert r.status_code == 422
    assert r.json()["detail"]["kode"] == "VALIDASI_ISIAN"


@pytest.mark.parametrize("path", [API_TRX, API_TGH])
def test_FR_LAP_04_OQ_38_ekspor_rentang_terbalik_ditolak(klien_admin, path):
    r = klien_admin.get(
        f"{path}/ekspor", params={"format": "pdf", "dari": "2026-10-02", "sampai": "2026-10-01"}
    )
    assert r.status_code == 422
    assert r.json()["detail"]["kode"] == "LAP_RENTANG_TIDAK_VALID"


def test_FR_LAP_04_ekspor_memanggil_query_yang_sama(db, monkeypatch, data_trx):
    """Satu sumber filter: ekspor tidak punya query sendiri."""
    dipanggil = []
    asli = laporan.laporan_transaksi

    def mata_mata(db_, f, **kw):
        dipanggil.append((f, kw))
        return asli(db_, f, **kw)

    monkeypatch.setattr(laporan, "laporan_transaksi", mata_mata)
    f = FilterTransaksi(status=StatusLaporan.TERLAMBAT)
    ekspor.ekspor_transaksi(db, f, "xlsx")
    assert dipanggil == [(f, {})]


# --------------------------------------------------------------------------- NFR-PRF-03


@pytest.fixture
def seribu_baris(db: Session) -> None:
    """1.000 item + 1.000 tagihan dibuat di test (seed performa tidak memuat transaksi)."""
    j = pabrik.judul(db, judul="Buku Performa Laporan Dengan Judul Agak Panjang")
    rak = pabrik.rak(db)
    a = pabrik.anggota(db, nama="Anggota Performa")
    admin = pabrik.admin(db)
    eks = db.scalars(
        insert(Eksemplar).returning(Eksemplar.id, sort_by_parameter_order=True),
        [{"judul_buku_id": j.id, "rak_id": rak.id, "status": "TERSEDIA"}] * 1_000,
    ).all()
    trx = db.scalars(
        insert(TransaksiPeminjaman).returning(TransaksiPeminjaman.id, sort_by_parameter_order=True),
        [
            {
                "anggota_id": a.id,
                "admin_id": admin.id,
                "tanggal_transaksi": date(2026, 9, 1),
                "status": "SELESAI",
            }
        ]
        * 1_000,
    ).all()
    item = db.scalars(
        insert(ItemTransaksi).returning(ItemTransaksi.id, sort_by_parameter_order=True),
        [
            {
                "transaksi_id": t,
                "eksemplar_id": e,
                "tanggal_pinjam": date(2026, 9, 1),
                "jatuh_tempo": date(2026, 10, 1),
                "tanggal_kembali": date(2026, 10, 9),
                "status": "DIKEMBALIKAN",
            }
            for t, e in zip(trx, eks, strict=True)
        ],
    ).all()
    db.execute(
        insert(Tagihan),
        [
            {
                "item_transaksi_id": i,
                "jenis": "DENDA",
                "nominal": 20_000,
                "status": "BELUM_LUNAS",
                "tanggal_dibentuk": date(2026, 10, 9),
            }
            for i in item
        ],
    )
    db.flush()
    # Statistik planner harus mencerminkan baris yang diukur. Suite me-rollback hampir semua insert,
    # sehingga autovacuum bisa mencatat tabel ini `relpages > 0, reltuples = 0` (halaman berisi
    # tuple mati saja); baris fixture yang belum di-commit tak terlihat autoanalyze, planner lalu
    # menaksir 1 baris dan memilih nested loop O(n^3): 98,9 dtk (transaksi) / 134,8 dtk (tagihan)
    # saat direproduksi. Di produksi data di-commit dan autoanalyze memperbaikinya <= 1 menit
    # (penyelidikan chore/selidiki-nfr-prf-03, progress.md 5.2.3/5.3.13). Di luar pengukuran waktu.
    db.execute(
        text(
            "ANALYZE item_transaksi, transaksi_peminjaman, eksemplar, anggota, judul_buku, "
            "tagihan, admin"
        )
    )


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
@pytest.mark.parametrize(("path", "kolom"), [(API_TRX, KOLOM_TRX), (API_TGH, KOLOM_TGH)])
def test_NFR_PRF_03_ekspor_1000_baris_maks_10_detik(
    klien_admin, seribu_baris, path, kolom, fmt, capsys
):
    mulai = time.perf_counter()
    r = klien_admin.get(f"{path}/ekspor", params={"format": fmt})
    durasi = time.perf_counter() - mulai
    assert r.status_code == 200, r.text
    with capsys.disabled():
        print(f"\nNFR-PRF-03 {path.rsplit('/', 1)[1]} {fmt}: {durasi:.2f} dtk")
    if fmt == "xlsx":
        assert len(_tabel_xlsx(r.content, kolom)[1]) == 1_000
    else:
        assert len(PdfReader(io.BytesIO(r.content)).pages) > 10
    assert durasi <= 10
